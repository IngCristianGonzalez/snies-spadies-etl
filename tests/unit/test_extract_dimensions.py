"""Tests unitarios para etl/extract_dimensions.py — Feature 016."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
import pytest

from etl.exceptions import DimensionTimeReadError
from etl.extract_dimensions import read_dim_tiempo


def _write_workbook(
    path: Path,
    sheets: dict[str, pd.DataFrame],
    *,
    preamble_rows: dict[str, list[list[str]] | None] | None = None,
) -> None:
    """Helper: write an Excel workbook with optional preamble rows before the header."""
    import openpyxl

    wb = openpyxl.Workbook()
    first = True
    for sheet_name, df in sheets.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)

        start_row = 1
        if preamble_rows and sheet_name in preamble_rows:
            for row in preamble_rows[sheet_name]:
                for col_idx, val in enumerate(row, start=1):
                    ws.cell(row=start_row, column=col_idx, value=val)
                start_row += 1

        for col_idx, col_name in enumerate(df.columns, start=1):
            ws.cell(row=start_row, column=col_idx, value=col_name)

        for row_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
            for col_idx, val in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=val)

    wb.save(path)


# ── T8: Archivo inexistente ────────────────────────────────────────────────


class TestFileNotFound:
    def test_raises_descriptive_error(self, tmp_path: Path) -> None:
        missing = tmp_path / "nonexistent.xlsx"
        with pytest.raises(DimensionTimeReadError, match="no existe"):
            read_dim_tiempo(missing)


# ── T9: Header después de filas introductorias y columnas adicionales ──────


class TestHeaderDetection:
    def test_header_after_preamble_rows(self, tmp_path: Path) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame(
            {
                "anio": [2020, 2021],
                "extra": ["z", "w"],
                "semestre": [1, 2],
            }
        )
        _write_workbook(
            path,
            {"Hoja1": data},
            preamble_rows={
                "Hoja1": [["NOT_DATA", "TITLE_ROW", "x"], ["NOT_DATA", "SUBTITLE", "y"]]
            },
        )
        result = read_dim_tiempo(path)
        assert list(result.columns) == ["anio", "semestre"]
        assert len(result) == 2
        assert result["anio"].tolist() == [2020, 2021]
        assert result["semestre"].tolist() == [1, 2]

    def test_selects_only_required_columns(self, tmp_path: Path) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame(
            {
                " descripcion ": ["A", "B"],
                " anio ": [2020, 2021],
                "extra_col": ["x", "y"],
                " semestre ": [1, 2],
            }
        )
        _write_workbook(path, {"Datos": data})
        result = read_dim_tiempo(path)
        assert list(result.columns) == ["anio", "semestre"]
        assert len(result) == 2


# ── T10: Ausencia de columnas y múltiples hojas candidatas ─────────────────


class TestAmbiguousSheets:
    def test_no_required_columns_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"col1": [1, 2], "col2": [3, 4]})
        _write_workbook(path, {"Hoja1": data})
        with pytest.raises(DimensionTimeReadError, match="No se encontraron"):
            read_dim_tiempo(path)

    def test_multiple_candidate_sheets_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        sheet1 = pd.DataFrame({"anio": [2020], "semestre": [1]})
        sheet2 = pd.DataFrame({"anio": [2021], "semestre": [2]})
        _write_workbook(path, {"Hoja1": sheet1, "Hoja2": sheet2})
        with pytest.raises(DimensionTimeReadError, match="varias hojas candidatas"):
            read_dim_tiempo(path)


# ── T11: Conversión de enteros y rechazo de valores inválidos ──────────────


class TestIntegerConversion:
    def test_textual_integers_converted(self, tmp_path: Path) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": ["2020", "2021"], "semestre": ["1", "2"]})
        _write_workbook(path, {"Hoja1": data})
        result = read_dim_tiempo(path)
        assert result["anio"].dtype == "int64"
        assert result["semestre"].dtype == "int64"

    def test_non_integer_values_omitted(self, tmp_path: Path, caplog) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": [2020, "abc", 2022], "semestre": [1, 2, "xyz"]})
        _write_workbook(path, {"Hoja1": data})
        with caplog.at_level(logging.WARNING):
            result = read_dim_tiempo(path)
        assert len(result) == 1
        assert result["anio"].tolist() == [2020]
        assert "row_omitted" in caplog.text or "omit" in caplog.text

    def test_negative_year_omitted(self, tmp_path: Path, caplog) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": [-1, 2021], "semestre": [1, 2]})
        _write_workbook(path, {"Hoja1": data})
        with caplog.at_level(logging.WARNING):
            result = read_dim_tiempo(path)
        assert len(result) == 1
        assert result["anio"].tolist() == [2021]

    def test_zero_year_omitted(self, tmp_path: Path) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": [0, 2021], "semestre": [1, 2]})
        _write_workbook(path, {"Hoja1": data})
        result = read_dim_tiempo(path)
        assert len(result) == 1

    def test_invalid_semestre_omitted(self, tmp_path: Path) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": [2020, 2021], "semestre": [3, 2]})
        _write_workbook(path, {"Hoja1": data})
        result = read_dim_tiempo(path)
        assert len(result) == 1
        assert result["semestre"].tolist() == [2]


# ── T12: Omisión de filas vacías y logging ─────────────────────────────────


class TestEmptyRows:
    def test_empty_rows_omitted_with_logging(self, tmp_path: Path, caplog) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": [2020, None, 2022], "semestre": [1, None, 2]})
        _write_workbook(path, {"Hoja1": data})
        with caplog.at_level(logging.WARNING):
            result = read_dim_tiempo(path)
        assert len(result) == 2
        assert "fila vacía" in caplog.text or "empty" in caplog.text


# ── T13: Conservación de primera ocurrencia y logging de duplicados ────────


class TestDeduplication:
    def test_first_occurrence_kept(self, tmp_path: Path) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": [2020, 2020, 2021], "semestre": [1, 1, 2]})
        _write_workbook(path, {"Hoja1": data})
        result = read_dim_tiempo(path)
        assert len(result) == 2
        assert result["anio"].tolist() == [2020, 2021]

    def test_duplicates_logged(self, tmp_path: Path, caplog) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": [2020, 2020, 2020], "semestre": [1, 1, 1]})
        _write_workbook(path, {"Hoja1": data})
        with caplog.at_level(logging.WARNING):
            result = read_dim_tiempo(path)
        assert len(result) == 1
        assert "duplicado" in caplog.text or "duplicate" in caplog.text


# ── T14: DataFrame vacío cuando no quedan filas válidas ────────────────────


class TestEmptyResult:
    def test_all_invalid_returns_empty_dataframe(self, tmp_path: Path, caplog) -> None:
        path = tmp_path / "dim_tiempo.xlsx"
        data = pd.DataFrame({"anio": ["abc", "def"], "semestre": ["xyz", "uvw"]})
        _write_workbook(path, {"Hoja1": data})
        with caplog.at_level(logging.ERROR):
            result = read_dim_tiempo(path)
        assert result.empty
        assert list(result.columns) == ["anio", "semestre"]
        assert "empty_result" in caplog.text or "vacío" in caplog.text
